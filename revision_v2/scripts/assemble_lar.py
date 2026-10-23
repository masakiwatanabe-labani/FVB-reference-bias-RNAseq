#!/usr/bin/env python3
"""Phase H-3: assemble the LAR/BMC-format main text and Additional file 1 from the section files, then build the docx files with
pandoc (--reference-doc = the v1 submission docx, so headings/body styles match).
  main text : title page (v1) -> structured Abstract (abstract_structured.md) -> Keywords -> Background (introduction.md, heading
              renamed) -> Methods -> Results -> Discussion -> Conclusions (split from discussion.md) -> List of abbreviations
              (abbreviations.md) -> Declarations (declarations.md) -> References (references.md, numbered by number_citations.py)
              -> Figure legends (main figures only) -> Table 1 (table1.md; footnote without script names)
  Additional file 1: Supplementary figure legends (S1–S3), Supplementary table captions, Supplementary Notes 1–2
Transformations applied to every section: local paths -> "the repository (file)"; numbered-list steps -> "Step n." paragraphs;
"* p < 0.05; ** p < 0.01; *** p < 0.001" asterisks escaped; HTML comments dropped; Markdown '---' rules dropped.
Outputs: revision_v2/submission/LAR_main_text_v2.md/.docx, LAR_additional_file_1.md/.docx (+ plain-text dumps for checks).
Usage: assemble_lar.py"""
import re, subprocess, os
PROJ="/usr/local/jupyter/FVB_bias"; M=f"{PROJ}/manuscript/revised_v2"; OUT=f"{PROJ}/revision_v2/submission"; REFDOC=f"{PROJ}/LAR text submit.docx"
os.makedirs(OUT, exist_ok=True)
def rd(f): return open(f"{M}/{f}", encoding="utf-8").read()
def clean(s):
    s=re.sub(r"<!--.*?-->", "", s, flags=re.S)
    s=re.sub(r"\n-{3,}\n", "\n", s)
    # local paths -> repository file names
    s=re.sub(r"`(?:revision_v2|analysis_v2|revision)/(?:[\w.\-]+/)*([\w.\-]+)`", r"the repository file `\1`", s)
    s=re.sub(r"\((?:revision_v2|analysis_v2)/(?:[\w.\-]+/)*([\w.\-]+)\)", r"(repository file \1)", s)
    s=s.replace("the repository file `manifest.tsv`","the sample manifest in the repository (`manifest.tsv`)")
    # numbered procedure lists -> Step paragraphs (Markdown "1. **X**: ..." at line start)
    s=re.sub(r"^(\d+)\.\s+\*\*([^*]+)\*\*:\s*", lambda m: f"Step {m.group(1)}. **{m.group(2)}**: ", s, flags=re.M)
    s=re.sub(r"^(\d+)\.\s+\*\*([^*]+)\*\*\s", lambda m: f"Step {m.group(1)}. **{m.group(2)}** ", s, flags=re.M)
    s=re.sub(r"\n(?!\n)(Step \d+\. )", r"\n\n\1", s)                      # each Step is its own paragraph (no soft line breaks)
    s=re.sub(r"`(?:revision_v2|analysis_v2)/(?:[\w.\-]+/)*`", "the repository", s)   # directory-only mentions
    s=re.sub(r"(?<![`\w/])(?:revision_v2|analysis_v2)/(?:[\w.\-]+/)*([\w.\-]+)?", lambda m: f"the repository file {m.group(1)}" if m.group(1) else "the repository", s)   # bare paths
    # significance asterisks (escape so pandoc does not treat them as emphasis)
    s=s.replace("* p < 0.05; ** p < 0.01; *** p < 0.001", "\\* p < 0.05; \\*\\* p < 0.01; \\*\\*\\* p < 0.001")
    # Phase S-1: command text keeps ASCII characters under pandoc's smart extension (which stays on for the prose dashes and
    # quotes): escape option dashes, quotes right after '=' (FILTER=".") and the quotes of a samtools -e expression; code spans
    # are verbatim already and are left alone
    def ascii_cmd(t):
        t=re.sub(r"(?<![\w-])--(?=[A-Za-z])", r"\\-\\-", t)
        t=re.sub(r'="([^"\s]*)"', r'=\\"\1\\"', t)
        return re.sub(r"(samtools view -e )'([^'\n]*)'", r"\1\\'\2\\'", t)
    s="".join(x if x.startswith("`") else ascii_cmd(x) for x in re.split(r"(`[^`\n]*`)", s))
    return s.strip()+"\n"
def heading_shift(s, top):  # make the section's own H1 the given top-level heading text
    return re.sub(r"^# .*$", f"# {top}", s, count=1, flags=re.M)
title=open(f"{OUT}/title_page.md", encoding="utf-8").read()
abstract=rd("abstract_structured.md"); keywords=rd("keywords.md")
background=heading_shift(rd("introduction.md"), "Background"); methods=rd("methods.md"); results=rd("results.md")
disc=rd("discussion.md"); i=disc.find("## Conclusions"); discussion=disc[:i]; conclusions="# Conclusions\n"+disc[i+len("## Conclusions"):]
abbrev=rd("abbreviations.md"); decl=rd("declarations.md"); refs=rd("references.md")
legends=rd("figure_legends.md"); j=legends.find("**Supplementary Figure S1"); main_legends=heading_shift(legends[:j], "Figure legends"); supp_legends="# Supplementary figure legends\n\n"+legends[j:]
table1=rd("table1.md"); table1=re.sub(r"^Generated from the sample manifest.*\n", "", table1, flags=re.M)
main="\n\n".join(clean(x) for x in (title, abstract, keywords, background, methods, results, discussion, conclusions, abbrev, decl, refs, main_legends, table1))
paper_title=[l.strip() for l in title.splitlines() if l.strip() and not l.startswith("**") and not l.startswith("#")][0]   # first plain line of the title page = the title
add="\n\n".join(clean(x) for x in ("# Additional file 1\n\nSupplementary figure legends, supplementary table captions and Supplementary Notes 1–2 for: "+paper_title+"\n", supp_legends, rd("supplementary_table_captions.md"), rd("supplementary_note1.md"), rd("supplementary_note2.md")))
open(f"{OUT}/LAR_main_text_v2.md","w",encoding="utf-8").write(main); open(f"{OUT}/LAR_additional_file_1.md","w",encoding="utf-8").write(add)
def fix_tables(path, font_pt=8, landscape_from=None):
    """Phase Q-5: pandoc tables ("Normal Table", equal grid columns, width 100 %) collapse when the docx is converted to PDF by
    LibreOffice. Give every table explicit widths in dxa that fit the text width (each column at least as wide as its longest word,
    the rest distributed by the amount of text), the "Table Grid" style, 8-pt text and a header row repeated on each page. Cell
    paragraphs lose pandoc's "Compact" style, which the localised reference docx lacks and which makes LibreOffice collapse the table
    (each cell rendered as a separate line, the preceding heading squeezed into a narrow column)."""
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import copy
    d=Document(path)
    if landscape_from:   # put the content from the heading landscape_from to the end (Table 1) in its own landscape section
        body=d.element.body; heads=[p for p in d.paragraphs if p.text.strip()==landscape_from]
        if heads:
            prev=heads[0]._p.getprevious()
            while prev is not None and prev.tag!=qn("w:p"): prev=prev.getprevious()
            final=body.find(qn("w:sectPr")); portrait=copy.deepcopy(final)
            pPr=prev.find(qn("w:pPr"))
            if pPr is None: pPr=OxmlElement("w:pPr"); prev.insert(0,pPr)
            pPr.append(portrait)
            pg=final.find(qn("w:pgSz")); w_,h_=pg.get(qn("w:w")),pg.get(qn("w:h")); pg.set(qn("w:w"),h_); pg.set(qn("w:h"),w_); pg.set(qn("w:orient"),"landscape")
    sec=d.sections[-1]; textw=int((sec.page_width-sec.left_margin-sec.right_margin)/635)   # EMU -> twips (the section of the tables)
    CH=int(font_pt*20*0.62)+1; PAD=200                                                                       # twips per character (bold header) at font_pt; cell padding
    grid_id=d.styles["Table Grid"].style_id                                                                  # the reference docx is localised: style ids are not the English names
    for t in d.tables:
        rows=[[c.text for c in r.cells] for r in t.rows]; n=len(rows[0])
        longest=[max(len(w) for r in rows for w in (r[j].split() or [""])) for j in range(n)]
        total=[sum(len(r[j]) for r in rows) for j in range(n)]
        base=[l*CH+PAD for l in longest]
        if sum(base)>textw: base=[int(b*textw/sum(base)) for b in base]   # never wider than the text width
        free=max(0,textw-sum(base)); tw=sum(total) or 1
        width=[b+int(free*x/tw) for b,x in zip(base,total)]; width[-1]+=textw-sum(width)
        tblPr=t._tbl.tblPr   # rebuilt in schema order (tblStyle, tblW, tblLayout, tblCellMar, tblLook); pandoc's order confuses LibreOffice
        look=tblPr.find(qn("w:tblLook"))
        for e in list(tblPr): tblPr.remove(e)
        st=OxmlElement("w:tblStyle"); st.set(qn("w:val"),grid_id); tblPr.append(st)
        tw_=OxmlElement("w:tblW"); tw_.set(qn("w:w"),str(textw)); tw_.set(qn("w:type"),"dxa"); tblPr.append(tw_)
        lay=OxmlElement("w:tblLayout"); lay.set(qn("w:type"),"fixed"); tblPr.append(lay)
        mar=OxmlElement("w:tblCellMar")
        for side,v in (("top",20),("left",60),("bottom",20),("right",60)):
            e=OxmlElement(f"w:{side}"); e.set(qn("w:w"),str(v)); e.set(qn("w:type"),"dxa"); mar.append(e)
        tblPr.append(mar)
        if look is not None: tblPr.append(look)
        grid=t._tbl.tblGrid
        for g in list(grid): grid.remove(g)
        for w in width:
            gc=OxmlElement("w:gridCol"); gc.set(qn("w:w"),str(w)); grid.append(gc)
        for i,row in enumerate(t.rows):
            if i==0:
                trPr=row._tr.get_or_add_trPr()
                for e in trPr.findall(qn("w:tblHeader")): trPr.remove(e)
                h=OxmlElement("w:tblHeader"); trPr.append(h)
            for j,cell in enumerate(row.cells):
                tcPr=cell._tc.get_or_add_tcPr()
                for e in tcPr.findall(qn("w:tcW")): tcPr.remove(e)
                cw=OxmlElement("w:tcW"); cw.set(qn("w:w"),str(width[j])); cw.set(qn("w:type"),"dxa"); tcPr.insert(0,cw)
                for p in cell.paragraphs:
                    ps=p._p.pPr.find(qn("w:pStyle")) if p._p.pPr is not None else None
                    if ps is not None and ps.get(qn("w:val")) not in {x.style_id for x in d.styles}: p._p.pPr.remove(ps)   # pandoc's "Compact" is absent from the reference docx; LibreOffice then collapses the table
                    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
                    for run in p.runs:
                        run.font.size=Pt(font_pt)
                        if i==0: run.font.bold=True
    d.save(path); return len(d.tables)
for name in ("LAR_main_text_v2","LAR_additional_file_1"):
    subprocess.run(["pandoc",f"{OUT}/{name}.md","-f","markdown","-t","docx","--reference-doc",REFDOC,"-o",f"{OUT}/{name}.docx"],check=True)
    print(name, "tables fixed for PDF conversion:", fix_tables(f"{OUT}/{name}.docx", landscape_from="Table 1. RNA-seq datasets used in this study" if name=="LAR_main_text_v2" else None))
    subprocess.run(["pandoc",f"{OUT}/{name}.docx","-t","plain","--wrap=none","-o",f"{OUT}/{name}.plain.txt"],check=True)
    print(name, "words:", len(open(f"{OUT}/{name}.md",encoding="utf-8").read().split()))
