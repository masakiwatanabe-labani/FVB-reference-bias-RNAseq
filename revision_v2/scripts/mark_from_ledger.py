#!/usr/bin/env python3
"""Phase J-6: blue-marked copy of the main text from the change ledger. Sentences whose ledger status is EDIT or NEW are blue
(RGB 0,0,255); KEEP sentences stay black; DROPped v1 sentences are not shown. Table 1 (v2 table) is blue. Legend line at the top.
Usage: mark_from_ledger.py [in.docx] [out.docx]  (defaults: revision_v2/submission/LAR_main_text_v2.docx -> *_marked.docx)
Prints the blue fraction (characters) per Heading-1 section."""
import re, sys, csv, copy, collections
from docx import Document
from docx.shared import RGBColor
from docx.text.run import Run
P="/usr/local/jupyter/FVB_bias"
IN=sys.argv[1] if len(sys.argv)>1 else f"{P}/revision_v2/submission/LAR_main_text_v2.docx"; OUT=sys.argv[2] if len(sys.argv)>2 else IN.replace(".docx","_marked.docx")
BLUE=RGBColor(0,0,255)
def norm(s):
    s=re.sub(r"\[CITE[^\]]*\]","",s); s=re.sub(r"\[\d+(?:[–,\-]\d+)*\]","",s); s=s.replace("\\","").replace("−","-").replace("–","-").replace("--","-")
    return re.sub(r"[\W_]+","",s).lower()
ABBR=("Fig","Figs","al","vs","e.g","i.e","ca","no","approx","Dr","St")
def sentences(text):
    spans=[]; start=0
    for m in re.finditer(r"[.!?;:](?=\s+\S|$)", text):
        before=text[:m.start()]
        if any(before.endswith(a) for a in ABBR): continue
        if text[m.start()]==":" and not text[m.end():].strip().startswith(("bias =","alt_count","Δ =","residual","Naive","FVB","samtools")): continue
        end=m.end(); spans.append((start,end)); start=end
    if start<len(text): spans.append((start,len(text)))
    return [(a,b) for a,b in spans if text[a:b].strip()]
V1N=set()
for line in open(f"{P}/revision_v2/submission/LAR_v1_text.md", encoding="utf-8"):
    for a,b in sentences(line.strip()): V1N.add(norm(line.strip()[a:b]))
    for piece in line.strip().split("\\"): V1N.add(norm(piece))
status={}
for r in csv.DictReader(open(f"{P}/revision_v2/change_ledger.tsv"), delimiter="\t"):
    if r["status"]!="DROP": status[norm(r["new_text"])]=r["status"]
def color_runs(par, mask):
    pos=0
    for r in list(par.runs):
        t=r.text; seg=mask[pos:pos+len(t)]; pos+=len(t)
        if not t: continue
        if all(seg): r.font.color.rgb=BLUE; continue
        if not any(seg): continue
        pieces=[]; i=0
        while i<len(t):
            j=i
            while j<len(t) and seg[j]==seg[i]: j+=1
            pieces.append((t[i:j],seg[i])); i=j
        r.text=pieces[0][0]
        if pieces[0][1]: r.font.color.rgb=BLUE
        prev=r._r
        for txt,col in pieces[1:]:
            nr=copy.deepcopy(r._r); prev.addnext(nr); run=Run(nr,par); run.text=txt; run.font.color.rgb=BLUE if col else None; prev=nr
def mark_all(par):
    for r in par.runs: r.font.color.rgb=BLUE
d=Document(IN); cur="(front)"; stat=collections.defaultdict(lambda:[0,0]); unknown=0; total=0
for p in d.paragraphs:
    if p.style.name=="Heading 1": cur=p.text.strip(); continue
    t=p.text
    if not t.strip() or cur=="References": continue
    mask=[False]*len(t)
    for a,b in sentences(t):
        k=norm(t[a:b]); total+=1
        if len(k)<4: continue
        st=status.get(k)
        if st is None: unknown+=1; st="KEEP" if (k in V1N or any(norm(x) in V1N for x in t[a:b].split("\n"))) else "NEW"   # outside the ledger (title page, legend titles, abbreviations): direct v1 comparison
        if "\n" in t[a:b] and cur not in ("(front)",):   # line-broken list (abbreviations): mark line by line against v1
            pos=a
            for x in t[a:b].split("\n"):
                if norm(x) and norm(x) not in V1N:
                    for i in range(pos,pos+len(x)): mask[i]=True
                pos+=len(x)+1
            continue
        if st in ("EDIT","NEW") and cur not in ("(front)",):
            for i in range(a,b): mask[i]=True
    if any(mask): color_runs(p,mask)
    stat[cur][0]+=sum(mask); stat[cur][1]+=len(t)
for tb in d.tables:
    for row in tb.rows:
        for c in row.cells:
            for p in c.paragraphs: mark_all(p)
leg=d.paragraphs[0].insert_paragraph_before("Blue text = changed or added relative to the submitted version; unchanged sentences are black; references renumbered."); mark_all(leg)
d.save(OUT)
print(OUT); print(f"sentences {total}, not in ledger {unknown}")
for s,(b,n) in stat.items():
    if n: print(f"  {b/n:5.0%} blue  {s}")
